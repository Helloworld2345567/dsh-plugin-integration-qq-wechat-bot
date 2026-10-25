import base64
import importlib.util
import io
import os
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import Mock, patch

ROOT = Path(__file__).resolve().parents[1]


def load(name, file):
    spec = importlib.util.spec_from_file_location(name, file)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


worker = load('bot_worker', ROOT / 'python/worker.py')
image_api = load('bot_image_api', ROOT / 'skills/image-generation/scripts/image_api.py')
media_api = load('bot_media_analysis', ROOT / 'skills/media-analysis/scripts/media_api.py')


class WorkerTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(dir=ROOT / '.codex')
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def job(self, operation, source=None, **params):
        return dict(id='test', scope='scope', operation=operation, deadline=time.time()*1000+30000,
                    inputs=[dict(path=str(source), name=source.name, kind='file', mime='application/octet-stream')] if source else [],
                    parameters=params, output=str(self.root), endpoint=dict(baseUrl='https://example.invalid/v1', model='fixture', keyEnv='TEST_PROVIDER_KEY'))

    def test_skill_manifest_dispatch_without_worker_switch(self):
        skills = self.root / 'skills'
        folder = skills / 'fixture'
        (folder / 'scripts').mkdir(parents=True)
        (folder / 'scripts/worker.py').write_text("def execute(job):\n    return {'text': job['parameters']['value'], 'coverage': 'FIXTURE'}\n", encoding='utf8')
        manifest = {'skill': 'fixture', 'worker': 'scripts/worker.py', 'operations': [{'id': 'fixture_extract'}]}
        (folder / 'operations.json').write_text(json.dumps(manifest), encoding='utf8')
        declarations = worker.load_skill_operations(skills)
        job = self.job('fixture_extract', value='skill-owned-result')
        job['handler'] = 'skill'
        with patch.object(worker, 'SKILL_OPERATIONS', declarations):
            self.assertEqual(worker.execute(job)['text'], 'skill-owned-result')
            with self.assertRaisesRegex(worker.JobError, 'UNKNOWN_MEDIA_OPERATION'):
                worker.execute({**job, 'operation': 'undeclared'})
        manifest['worker'] = '../../outside.py'
        (folder / 'operations.json').write_text(json.dumps(manifest), encoding='utf8')
        with self.assertRaisesRegex(worker.JobError, 'MANIFEST_INVALID'):
            worker.load_skill_operations(skills)

    def test_shipped_skill_operations_dispatch_declared_entry(self):
        self.assertEqual(set(worker.SKILL_OPERATIONS), {'transcribe', 'video_analyze', 'image_download', 'video_download', 'voice_speak', 'voice_chant'})
        job = self.job('image_download')
        job['handler'] = 'skill'
        # Real entry import and error translation, without a provider/network call.
        job['inputs'] = [{ 'path': 'unused', 'name': 'unused', 'kind': 'image', 'mime': 'image/png'}] * 2
        with self.assertRaisesRegex(worker.JobError, 'ONE_MEDIA_INPUT_REQUIRED'):
            worker.execute(job)

    def test_unlimited_media_job_and_platform_source_routing(self):
        job = self.job('transcribe')
        job.update(timeoutMs=0, deadline=0)
        worker._validate_protocol(job)
        self.assertIsNone(media_api.remaining(job))
        from video_understanding_core.models import Limits
        self.assertIsNone(Limits(job_timeout=None, model_timeout=None).job_timeout)
        from video_understanding_core.sources import platform_source
        self.assertEqual(platform_source('https://v.douyin.com/34Ae2yNHPXs/')['kind'], 'douyin')
        self.assertEqual(platform_source('https://www.bilibili.com/video/BV1xx411c7mD?p=2')['page'], 2)

    def test_signed_download_preserves_escaped_url(self):
        import asyncio
        from video_understanding_core import network
        url = 'https://example.com/image%2Fname.webp?x-signature=abc%2Fdef%3D'
        class Response:
            status = 200
            headers = {}
            def __init__(self): self.content = self
            async def __aenter__(self): return self
            async def __aexit__(self, *args): pass
            async def iter_chunked(self, size): yield b'image-bytes'
        class Client:
            async def __aenter__(self): return self
            async def __aexit__(self, *args): pass
            def get(client, requested, **kwargs):
                self.assertEqual(str(requested), url)
                self.assertFalse(kwargs['allow_redirects'])
                return Response()
        target = self.root / 'download.bin'
        with patch.object(network.aiohttp, 'TCPConnector', return_value=object()), patch.object(network.aiohttp, 'ClientSession', return_value=Client()):
            self.assertEqual(asyncio.run(network.download(url, target, 100, None)), 11)
        self.assertEqual(target.read_bytes(), b'image-bytes')

    def test_public_link_parser_keeps_identity_and_gallery_originals(self):
        from platform_links import describe_link, normalize_result
        from video_understanding_core.models import VideoError
        source = describe_link('https://www.douyin.com/note/7662259898261030499')
        payload = {'succ': True, 'data': {'platform': '抖音', 'video_id': source['work_id'], 'image_list': ['https://example.com/original.webp', {'url': 'https://example.com/second.webp'}]}}
        self.assertEqual(len(normalize_result(payload, source)['images']), 2)
        payload['data']['video_id'] = 'different'
        with self.assertRaisesRegex(VideoError, 'SOURCE_IDENTITY_MISMATCH'):
            normalize_result(payload, source)
        with self.assertRaisesRegex(VideoError, 'PUBLIC_LINK_PARSE_FAILED'):
            normalize_result({'succ': False}, source)
        with self.assertRaises(VideoError):
            describe_link('https://example.com/not-a-platform')

    def test_original_image_selection_and_bytes(self):
        sys.path.insert(0, str(ROOT / 'skills/media-analysis/scripts'))
        from image_download import page_images, image_file
        html = '<img src="thumb.jpg" data-original="full.png"><img srcset="small.jpg 200w, large.jpg 1200w"><a href="original.webp"><img src="preview.jpg"></a>'
        self.assertEqual(page_images(html, 'https://example.com/post/'), ['https://example.com/post/full.png', 'https://example.com/post/large.jpg', 'https://example.com/post/original.webp'])
        from PIL import Image
        file = self.root / 'download.bin'
        Image.new('RGB', (31, 27), 'green').save(file, format='PNG')
        original = file.read_bytes()
        result = image_file(file, 1)
        self.assertEqual(result['mime'], 'image/png')
        self.assertEqual(Path(result['path']).read_bytes(), original)
        file.write_text('<html>not an image</html>')
        self.assertIsNone(image_file(file, 2))

    def test_documents_use_actual_contents_and_preserve_unicode(self):
        from docx import Document
        from openpyxl import Workbook
        text = self.root / 'text.txt'; text.write_text('你好，文档内容', encoding='utf-8')
        self.assertEqual(worker.execute(self.job('document_read', text))['text'], '你好，文档内容')
        csv = self.root / 'data.csv'; csv.write_text('name,value\n猫,7', encoding='utf-8')
        self.assertIn('猫,7', worker.execute(self.job('document_read', csv))['text'])
        doc = Document(); doc.add_paragraph('document fixture 42'); file = self.root / 'document.docx'; doc.save(file)
        self.assertIn('fixture 42', worker.execute(self.job('document_read', file))['text'])
        book = Workbook(); book.active.append(['known cell', 63]); file = self.root / 'sheet.xlsx'; book.save(file); book.close()
        self.assertIn('known cell\t63', worker.execute(self.job('document_read', file))['text'])

    def test_scanned_pdf_returns_pixels_for_native_chat_without_vision_service(self):
        from pypdf import PdfWriter
        file = self.root / 'scan.pdf'; writer = PdfWriter(); writer.add_blank_page(width=200, height=200)
        with file.open('wb') as handle: writer.write(handle)
        job = self.job('document_read', file); job['endpoint'] = None
        with patch('requests.Session') as request:
            result = worker.execute(job); request.assert_not_called()
        self.assertFalse(result['incomplete']); self.assertIn('当前聊天模型', result['text'])
        self.assertEqual(result['coverage'], 'DOCUMENT_TEXT_AND_PAGE_PIXELS')
        from PIL import Image
        with Image.open(result['contextAssets'][0]['path']) as page:
            self.assertEqual(page.format, 'JPEG'); self.assertGreater(page.width, 0)

    def test_scan_render_limit_reports_unread_pages(self):
        from pypdf import PdfWriter
        file=self.root/'many-pages.pdf'; writer=PdfWriter()
        for _ in range(11): writer.add_blank_page(width=200,height=200)
        with file.open('wb') as handle: writer.write(handle)
        job=self.job('document_read',file);job['endpoint']=None
        result=worker.execute(job);self.assertTrue(result['incomplete']);self.assertEqual(len(result['contextAssets']),10)

    def test_reference_edit_requires_real_input_and_canonical_client(self):
        with self.assertRaisesRegex(worker.JobError, 'REFERENCE_IMAGE_REQUIRED'): worker.execute(self.job('image_edit', prompt='edit'))
        from PIL import Image
        file = self.root / 'ref.png'; Image.new('RGB', (20, 20), 'green').save(file)
        job = self.job('image_edit', file, prompt='edit this'); job['inputs'][0]['kind'] = 'image'
        def fake(command, **kwargs):
            self.assertEqual(Path(command[1]), ROOT / 'skills/image-generation/scripts/image_api.py')
            self.assertIn('--image', command); self.assertIn(str(file), command)
            self.assertEqual(command[command.index('--api-key-env')+1], 'TEST_PROVIDER_KEY')
            output = self.root / 'generated.jpg'; Image.new('RGB',(10,10),'blue').save(output)
            (self.root / 'image-result.json').write_text(json.dumps({'outputs':[str(output)]}),encoding='utf-8')
            return Mock(returncode=0)
        with patch.object(worker.subprocess, 'run', side_effect=fake): self.assertEqual(len(worker.execute(job)['assets']), 1)

    def test_image_client_normalizes_jpeg_and_does_not_replay_paid_requests(self):
        from PIL import Image
        buf = io.BytesIO(); Image.new('RGBA', (12,12), (0,255,0,100)).save(buf, 'PNG')
        files = image_api.save_images([(buf.getvalue(),'image/png')],self.root,'fixture')
        self.assertEqual(len(files),1)
        with Image.open(files[0]) as result: self.assertEqual(result.format,'JPEG')
        session = Mock(); session.request.side_effect=TimeoutError('uncertain upstream outcome')
        with self.assertRaises(TimeoutError): image_api.request_with_retry(session,'POST','https://example.invalid',headers={},timeout=1)
        self.assertEqual(session.request.call_count,1)

    def test_shared_image_cli_inherits_extension_variables(self):
        with patch.dict(os.environ,dict(BOT_IMAGE_BASE_URL='https://example.com/v1',BOT_IMAGE_MODEL='custom-model',BOT_IMAGE_KEY_ENV='CUSTOM_IMAGE_KEY'),clear=True),patch.object(sys,'argv',['image_api.py','generate','--prompt','test','--dry-run']):
            args=image_api.parse_args();self.assertEqual(args.base_url,'https://example.com/v1');self.assertEqual(args.model,'custom-model');self.assertEqual(args.api_key_env,'CUSTOM_IMAGE_KEY')

    @unittest.skipUnless(shutil.which('ffmpeg') and shutil.which('ffprobe'), 'ffmpeg not installed')
    def test_audio_analysis_sends_actual_wav_and_detects_missing_media(self):
        import wave
        file=self.root/'audio.wav'
        with wave.open(str(file),'wb') as wav:
            wav.setnchannels(1);wav.setsampwidth(2);wav.setframerate(16000);wav.writeframes(b'\0\0'*1600)
        job=self.job('transcribe',file,question='转写');job['endpoint']['protocol']='chat_completions'
        def response(body):return {'choices':[{'finish_reason':'stop','message':{'content':json.dumps(body)}}]}
        value={'received_media':True,'transcript':'known phrase','summary':'测试','uncertainty':[]}
        with patch.object(media_api,'request',return_value=response(value)) as request:
            result=media_api.execute(job);self.assertEqual(json.loads(result['text'])['transcript'],'known phrase')
            payload=request.call_args.kwargs['body'];self.assertEqual(payload['model'],'fixture');self.assertEqual(request.call_args.args[1],'/chat/completions')
            audio=payload['messages'][0]['content'][1]['input_audio'];self.assertEqual(audio['format'],'wav');self.assertTrue(base64.b64decode(audio['data']).startswith(b'RIFF'))
            self.assertNotIn('TEST_PROVIDER_KEY',json.dumps(payload));self.assertGreater(json.loads(result['text'])['coverage']['endSeconds'],0)
        with patch.object(media_api,'request',return_value=response({'received_media':False})):
            with self.assertRaisesRegex(media_api.AnalysisError,'MEDIA_NOT_RECEIVED'):media_api.execute(job)
        job['endpoint']['protocol']='transcriptions'
        with patch.object(media_api,'request',return_value={'text':'known phrase'}) as request:
            self.assertEqual(json.loads(media_api.execute(job)['text'])['transcript'],'known phrase');self.assertIn('file',request.call_args.kwargs['files'])
        with patch.object(media_api,'request',return_value={'text':''}):
            with self.assertRaisesRegex(media_api.AnalysisError,'EMPTY_TRANSCRIPT'):media_api.execute(job)

    def test_analysis_http_uses_only_injected_key_and_does_not_retry_uncertain_requests(self):
        import requests
        session=Mock();session.__enter__=Mock(return_value=session);session.__exit__=Mock(return_value=False)
        response=Mock();response.__enter__=Mock(return_value=response);response.__exit__=Mock(return_value=False)
        response.status_code=200;response.iter_content.return_value=[b'{"text":"ok"}'];session.post.return_value=response
        with patch.dict(os.environ,{'TEST_PROVIDER_KEY':'fixture-secret'},clear=True),patch('requests.Session',return_value=session):
            result=media_api.request(self.job('transcribe'),'/chat/completions',body={'model':'fixture'})
            self.assertEqual(result['text'],'ok');self.assertEqual(session.post.call_args.kwargs['headers']['Authorization'],'Bearer fixture-secret')
            self.assertFalse(session.post.call_args.kwargs['allow_redirects']);self.assertFalse(session.trust_env)
            session.post.reset_mock();session.post.side_effect=requests.Timeout('secret must not appear')
            with self.assertRaisesRegex(media_api.AnalysisError,'MODEL_TIMEOUT_MAY_BE_BILLED'):media_api.request(self.job('transcribe'),'/chat/completions')
            self.assertEqual(session.post.call_count,1)

    def test_usage_journal_survives_failure_and_captures_actual_model_without_content(self):
        import requests
        from model_usage import journal
        job = self.job('transcribe'); job['usagePath'] = str(self.root/'usage.json')
        session=Mock();session.__enter__=Mock(return_value=session);session.__exit__=Mock(return_value=False)
        session.post.side_effect=requests.Timeout('fixture')
        with patch.dict(os.environ,{'TEST_PROVIDER_KEY':'fixture-secret'},clear=True),patch('requests.Session',return_value=session):
            with self.assertRaises(media_api.AnalysisError): media_api.request(job,'/chat/completions')
        records=json.loads(Path(job['usagePath']).read_text()); self.assertEqual(len(records),1)
        self.assertTrue(records[0]['attempted']);self.assertNotIn('usage',records[0]);self.assertNotIn('fixture-secret',json.dumps(records))
        journal(job, records[0], {'model':'actual-model','usage':{'prompt_tokens':10,'completion_tokens':5,'completion_tokens_details':{'reasoning_tokens':3,'text':'secret'},'content':'secret'},'text':'private'})
        records=json.loads(Path(job['usagePath']).read_text());self.assertEqual(len(records),1);self.assertEqual(records[0]['model'],'actual-model')
        self.assertEqual(records[0]['usage']['completion_tokens_details'],{'reasoning_tokens':3});self.assertNotIn('private',json.dumps(records));self.assertNotIn('secret',json.dumps(records))

    def test_local_processing_and_validation_failures_do_not_claim_model_calls(self):
        file=self.root/'test.txt';file.write_text('fixture')
        job=self.job('document_read',file);job['usagePath']=str(self.root/'usage.json')
        worker.execute(job);self.assertFalse(Path(job['usagePath']).exists())
        with self.assertRaises(worker.JobError):worker.execute(self.job('image_edit'))
        self.assertFalse((self.root/'model-usage.json').exists())

    @unittest.skipUnless(shutil.which('ffmpeg'), 'ffmpeg not installed')
    def test_video_payload_contains_actual_mp4_and_checks_reported_media(self):
        from video_understanding_core.backend import BackendConfig,ChatCompletionsGemini,validate_result
        from video_understanding_core.models import Limits,MediaProbe,PreparedMedia,VideoError
        source=self.root/'clip.mp4'
        subprocess.run(['ffmpeg','-v','error','-f','lavfi','-i','color=c=red:s=96x96:d=1','-c:v','libx264',str(source)],check=True,capture_output=True)
        prepared=PreparedMedia(source,MediaProbe(1,True,False,source.stat().st_size),1,0,1)
        backend=ChatCompletionsGemini(BackendConfig('https://example.invalid/v1','test-secret','fixture'),Limits())
        payload=json.loads(backend.payload(prepared,'画面是什么？'))
        data=payload['messages'][0]['content'][1]['image_url']['url'];self.assertTrue(data.startswith('data:video/mp4;base64,'));self.assertEqual(base64.b64decode(data.split(',',1)[1]),source.read_bytes())
        with self.assertRaisesRegex(VideoError,'MEDIA_NOT_RECEIVED'):validate_result('{"received_media":false}',prepared)
        value={'received_media':True,'summary':'red','visual_observations':[{'at_seconds':0,'text':'red'}],'audio_observations':[{'at_seconds':0,'text':'fabricated'}],'spoken_quotes':[],'onscreen_text':[],'inferences':[],'uncertainty':[],'warnings':[]}
        self.assertEqual(validate_result(json.dumps(value),prepared)['audio_observations'],[])

    @unittest.skipUnless(shutil.which('ffmpeg'), 'ffmpeg not installed')
    def test_real_fixed_ffmpeg_conversion(self):
        source=self.root/'clip.mp4'
        subprocess.run(['ffmpeg','-v','error','-f','lavfi','-i','color=c=red:s=96x96:d=1','-c:v','libx264',str(source)],check=True,capture_output=True)
        result=worker.execute(self.job('convert_gif',source,durationSeconds=0.5))
        from PIL import Image
        with Image.open(result['assets'][0]['path']) as image: self.assertEqual(image.format,'GIF')
        with self.assertRaisesRegex(worker.JobError,'INVALID_MEDIA_RANGE'):worker.execute(self.job('convert_mp4',source,durationSeconds=999))

    def test_worker_protocol_rejects_arbitrary_operation_and_expired_jobs(self):
        for job,code in [(self.job('run_script'),'UNKNOWN_MEDIA_OPERATION'),(self.job('speak'),'UNKNOWN_MEDIA_OPERATION'),({**self.job('transcribe'),'deadline':0},'JOB_EXPIRED')]:
            result=subprocess.run([sys.executable,str(ROOT/'python/worker.py')],input=json.dumps(job),capture_output=True,text=True)
            self.assertEqual(result.returncode,1);self.assertEqual(json.loads(result.stdout)['error']['code'],code)


if __name__ == '__main__': unittest.main()
