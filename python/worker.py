"""One fixed media job over stdin/stdout. No script, import-name or shell API."""
from __future__ import annotations
import asyncio
import contextlib
import io
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import time
import zipfile

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent / "vendor"))
PROTOCOL = "dsh.media-worker"
PROTOCOL_VERSION = 1


class JobError(Exception):
    pass


def checked_number(value, minimum, maximum):
    if type(value) not in (int, float) or not math.isfinite(value) or not minimum <= value <= maximum:
        raise JobError("INVALID_MEDIA_RANGE")
    return value


def _validate_json_shape(value, *, depth=0):
    """Bound the shape before a handler can hand parameters to a provider."""
    if depth > 6:
        raise JobError("MEDIA_WORKER_PARAMETERS_INVALID")
    if isinstance(value, str):
        if len(value) > 100_000:
            raise JobError("MEDIA_WORKER_PARAMETERS_INVALID")
    elif isinstance(value, list):
        if len(value) > 64:
            raise JobError("MEDIA_WORKER_PARAMETERS_INVALID")
        for item in value:
            _validate_json_shape(item, depth=depth + 1)
    elif isinstance(value, dict):
        if len(value) > 64:
            raise JobError("MEDIA_WORKER_PARAMETERS_INVALID")
        for key, item in value.items():
            if not isinstance(key, str) or not 1 <= len(key) <= 128:
                raise JobError("MEDIA_WORKER_PARAMETERS_INVALID")
            _validate_json_shape(item, depth=depth + 1)
    elif value is not None and not isinstance(value, (bool, int, float)):
        raise JobError("MEDIA_WORKER_PARAMETERS_INVALID")


def image(job):
    params = job["parameters"]
    operation = "edit" if job["operation"] == "image_edit" else "generate"
    if operation == "edit" and not job["inputs"]:
        raise JobError("REFERENCE_IMAGE_REQUIRED")
    endpoint = job["endpoint"]
    record = Path(job["output"]) / "image-result.json"
    command = [sys.executable, str(ROOT / "skills/image-generation/scripts/image_api.py"), operation,
               "--prompt", str(params.get("prompt", "")), "--model", endpoint["model"],
               "--base-url", endpoint["baseUrl"], "--api-key-env", endpoint["keyEnv"],
               "--output-dir", job["output"], "--output-format", "jpeg", "--n", "1",
               "--timeout", str(job.get("timeoutMs", 0) / 1000), "--record", str(record)]
    if job.get('usagePath'):
        command.extend(['--usage-record', job['usagePath']])
    if operation == "edit":
        for source in job["inputs"]:
            if source["kind"] != "image":
                raise JobError("REFERENCE_MUST_BE_IMAGE")
            command.extend(["--image", source["path"]])
    result = subprocess.run(command, capture_output=True, timeout=(job.get("timeoutMs", 0) / 1000 or None), check=False)
    if result.returncode:
        raise JobError("IMAGE_PROVIDER_FAILED")
    value = json.loads(record.read_text(encoding="utf-8"))
    # The canonical client records paths; never take a path from a model response.
    if "outputs" not in value:
        raise JobError("IMAGE_OUTPUT_RECORD_MISSING")
    files = value["outputs"]
    if not isinstance(files, list) or not files:
        raise JobError("IMAGE_OUTPUT_RECORD_INVALID")
    if not all(isinstance(p, str) and p.strip() for p in files):
        raise JobError("IMAGE_OUTPUT_RECORD_INVALID")
    if len(files) != 1:
        raise JobError("IMAGE_OUTPUT_COUNT_INVALID")
    return {"assets": [{"path": files[0], "name": "image.jpg", "mime": "image/jpeg", "kind": "image"}]}


def skill(job):
    import importlib.util
    entry = SKILL_OPERATIONS.get(job['operation'])
    if entry is None:
        raise JobError('UNKNOWN_MEDIA_OPERATION')
    script, _ = entry
    spec = importlib.util.spec_from_file_location('bot_skill_worker', script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    try:
        return module.execute(job)
    except Exception as error:
        raise JobError(getattr(error, 'code', type(error).__name__)) from None


def load_skill_operations(root):
    result = {}
    for manifest_path in sorted(root.glob('*/operations.json')):
        manifest = json.loads(manifest_path.read_text(encoding='utf8'))
        script = (manifest_path.parent / manifest['worker']).resolve()
        if manifest['skill'] != manifest_path.parent.name or not script.is_relative_to(manifest_path.parent.resolve()) or not script.is_file():
            raise JobError('SKILL_OPERATION_MANIFEST_INVALID')
        for operation in manifest['operations']:
            if operation['id'] in result:
                raise JobError('SKILL_OPERATION_MANIFEST_INVALID')
            result[operation['id']] = (script, operation)
    return result


SKILL_OPERATIONS = load_skill_operations(ROOT / 'skills')


def document(job):
    if len(job["inputs"]) != 1:
        raise JobError("ONE_DOCUMENT_REQUIRED")
    source = job["inputs"][0]
    file = Path(source["path"])
    suffix = Path(source["name"]).suffix.lower()
    incomplete = False
    context_assets = []
    if suffix in (".docx", ".xlsx"):
        with zipfile.ZipFile(file) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024**2:
                raise JobError("DOCUMENT_EXPANSION_LIMIT")
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(file)
        count = min(len(reader.pages), 100)
        texts = [page.extract_text() or "" for page in reader.pages[:count]]
        text = "\n".join(texts)
        missing = [i for i, value in enumerate(texts) if not value.strip()]
        if missing:
            import pypdfium2
            pdf = pypdfium2.PdfDocument(file)
            for i in missing[:10]:
                page = pdf[i]
                width, height = page.get_size()
                bitmap = page.render(scale=min(1.5, 1800 / max(width, height)))
                output = Path(job["output"]) / f"page-{i+1}.jpg"
                bitmap.to_pil().convert("RGB").save(output, "JPEG")
                texts[i] = f"[第{i+1}页：无可提取文字，已附真实页面图像，请由当前聊天模型读取。]"
                context_assets.append(dict(path=str(output), name=f"page-{i+1}.jpg", mime="image/jpeg", kind="image"))
                bitmap.close(); page.close()
            pdf.close(); text = "\n".join(texts)
        incomplete = len(reader.pages) > count or any(not text.strip() for text in texts)
    elif suffix == ".docx":
        from docx import Document
        doc = Document(file)
        text = "\n".join([p.text for p in doc.paragraphs] + ["\t".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(file, read_only=True, data_only=True)
        rows, cells = [], 0
        for sheet in book:
            rows.append(f"[{sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells += len(row)
                if cells > 10000:
                    incomplete = True; break
                rows.append("\t".join("" if value is None else str(value) for value in row))
            if incomplete: break
        book.close(); text = "\n".join(rows)
    elif suffix in (".txt", ".csv", ".tsv", ".md", ".json"):
        try: text = file.read_text(encoding="utf-8-sig")
        except UnicodeError: text = file.read_text(encoding="gb18030")
    else:
        raise JobError("UNSUPPORTED_DOCUMENT_FORMAT")
    if len(text) > 100000:
        text = text[:100000]; incomplete = True
    return {"text": text, "coverage": "DOCUMENT_TEXT_AND_PAGE_PIXELS" if context_assets else "DOCUMENT_EXTRACTION", "incomplete": incomplete, "contextAssets": context_assets}


def convert(job):
    if len(job["inputs"]) != 1:
        raise JobError("ONE_VIDEO_REQUIRED")
    params = job["parameters"]
    start = checked_number(params.get("startSeconds", 0), 0, 36000)
    duration = checked_number(params.get("durationSeconds", 6), 0.1, 30)
    output = Path(job["output"]) / "result.mp4"
    command = ["ffmpeg", "-v", "error", "-nostdin", "-ss", str(start), "-i", job["inputs"][0]["path"], "-t", str(duration)]
    command += ["-c:v", "libx264", "-c:a", "aac", "-movflags", "+faststart"]
    result = subprocess.run([*command, str(output)], capture_output=True, timeout=(job.get("timeoutMs", 0) / 1000 or None))
    if result.returncode or not output.exists():
        raise JobError("FFMPEG_FAILED")
    return {"assets": [{"path": str(output), "name": output.name, "mime": "video/mp4", "kind": "video"}]}


class OperationRegistry:
    """Versioned worker dispatch. Adding an operation is one registration, not a switch edit."""
    def __init__(self):
        self._handlers = {}

    def register(self, name, handler):
        if not isinstance(name, str) or not name or name in self._handlers:
            raise JobError("MEDIA_OPERATION_REGISTRATION_INVALID")
        self._handlers[name] = handler

    def execute(self, job):
        handler = self._handlers.get(job.get("handler"))
        if handler is None:
            raise JobError("UNKNOWN_MEDIA_OPERATION")
        return handler(job)


def _validate_result(result):
    if not isinstance(result, dict):
        raise JobError("MEDIA_WORKER_RESULT_INVALID")
    for key in ("text", "coverage"):
        if key in result and result[key] is not None and not isinstance(result[key], str):
            raise JobError("MEDIA_WORKER_RESULT_INVALID")
    if "incomplete" in result and not isinstance(result["incomplete"], bool):
        raise JobError("MEDIA_WORKER_RESULT_INVALID")
    for key in ("assets", "contextAssets"):
        files = result.get(key, [])
        if not isinstance(files, list) or len(files) > 16:
            raise JobError("MEDIA_WORKER_RESULT_INVALID")
        for file in files:
            if not isinstance(file, dict) or any(not isinstance(file.get(field), str) or not file[field] for field in ("path", "name", "mime", "kind")):
                raise JobError("MEDIA_WORKER_RESULT_INVALID")
    return result


def _validate_usage_journal(job):
    journal = Path(job["usagePath"])
    if not journal.exists():
        return
    try:
        value = json.loads(journal.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        raise JobError("MEDIA_USAGE_JOURNAL_INVALID") from None
    if not isinstance(value, list) or len(value) > 256:
        raise JobError("MEDIA_USAGE_JOURNAL_INVALID")
    for item in value:
        if not isinstance(item, dict) or not isinstance(item.get("id"), str) or not 1 <= len(item["id"]) <= 128 or item.get("attempted") is not True or not isinstance(item.get("at"), int):
            raise JobError("MEDIA_USAGE_JOURNAL_INVALID")
        if "model" in item and item["model"] is not None and (not isinstance(item["model"], str) or len(item["model"]) > 256):
            raise JobError("MEDIA_USAGE_JOURNAL_INVALID")
        if "usage" in item and item["usage"] is not None and not isinstance(item["usage"], dict):
            raise JobError("MEDIA_USAGE_JOURNAL_INVALID")


def _social(job):
    from social_posts import read_post
    return {"text": json.dumps(asyncio.run(read_post(str(job["parameters"].get("url", "")))), ensure_ascii=False), "coverage": "PAGE_TEXT_ONLY"}


def _validate_protocol(job):
    if not isinstance(job, dict):
        raise JobError("MEDIA_WORKER_PROTOCOL_UNSUPPORTED")
    if 'processingSlots' in job and (type(job['processingSlots']) is not int or not 1 <= job['processingSlots'] <= 2):
        raise JobError('MEDIA_WORKER_SLOTS_INVALID')
    # Keep the pre-Phase-5 direct Python API usable for existing integrations;
    # jobs crossing the process boundary are emitted with an explicit version.
    if "protocol" not in job and "protocolVersion" not in job:
        job["protocol"] = PROTOCOL
        job["protocolVersion"] = PROTOCOL_VERSION
        job.setdefault("timeoutMs", 180000)
        job.setdefault("quotaUnits", 1)
        job.setdefault("usagePath", str(Path(job.get("output", ".")) / "model-usage.json"))
        job.setdefault("handler", job.get("operation"))
    if "handler" not in job:
        job["handler"] = job.get("operation")
    if job.get("protocol") != PROTOCOL or job.get("protocolVersion") != PROTOCOL_VERSION:
        raise JobError("MEDIA_WORKER_PROTOCOL_UNSUPPORTED")
    if not isinstance(job.get("id"), str) or not 1 <= len(job["id"]) <= 128 or not isinstance(job.get("scope"), str) or not 1 <= len(job["scope"]) <= 256 or not isinstance(job.get("operation"), str) or not 2 <= len(job["operation"]) <= 81 or not isinstance(job.get("handler"), str) or not 2 <= len(job["handler"]) <= 81:
        raise JobError("MEDIA_WORKER_JOB_INVALID")
    if not isinstance(job.get("parameters"), dict):
        raise JobError("MEDIA_WORKER_PARAMETERS_INVALID")
    try:
        if len(json.dumps(job["parameters"], ensure_ascii=False).encode("utf-8")) > 256 * 1024:
            raise JobError("MEDIA_WORKER_PARAMETERS_TOO_LARGE")
        _validate_json_shape(job["parameters"])
    except (TypeError, ValueError):
        raise JobError("MEDIA_WORKER_PARAMETERS_INVALID") from None
    if not isinstance(job.get("output"), str) or not job["output"] or len(job["output"]) > 4096 or not isinstance(job.get("usagePath"), str) or not job["usagePath"] or len(job["usagePath"]) > 4096:
        raise JobError("MEDIA_WORKER_PATH_INVALID")
    if not isinstance(job.get("inputs"), list) or len(job["inputs"]) > 8:
        raise JobError("INVALID_INPUTS")
    for item in job["inputs"]:
        if not isinstance(item, dict) or not all(isinstance(item.get(key), str) and item[key] for key in ("path", "name", "mime", "kind")) or any(len(item[key]) > limit for key, limit in (("path", 4096), ("name", 512), ("mime", 256), ("kind", 32))):
            raise JobError("MEDIA_WORKER_INPUT_INVALID")
    endpoint = job.get("endpoint")
    if endpoint is not None and (not isinstance(endpoint, dict) or not isinstance(endpoint.get("model"), str) or not endpoint["model"] or len(endpoint["model"]) > 256 or not isinstance(endpoint.get("baseUrl"), str) or not endpoint["baseUrl"] or len(endpoint["baseUrl"]) > 2048 or not isinstance(endpoint.get("keyEnv"), str) or not endpoint["keyEnv"] or not endpoint["keyEnv"][0].isalpha() or not endpoint["keyEnv"].replace("_", "").isalnum() or not endpoint["keyEnv"].isupper()):
        raise JobError("MEDIA_WORKER_ENDPOINT_INVALID")
    if not isinstance(job.get("deadline"), (int, float)) or not math.isfinite(job["deadline"]) or (job.get("timeoutMs") != 0 and job["deadline"] < time.time() * 1000):
        raise JobError("JOB_EXPIRED")
    if not isinstance(job.get("timeoutMs"), int) or (job["timeoutMs"] != 0 and job["timeoutMs"] < 1000) or job["timeoutMs"] > 15 * 60 * 1000:
        raise JobError("MEDIA_WORKER_TIMEOUT_INVALID")
    if not isinstance(job.get("quotaUnits"), int) or job["quotaUnits"] < 1 or job["quotaUnits"] > 10000:
        raise JobError("MEDIA_WORKER_QUOTA_INVALID")


OPERATIONS = OperationRegistry()
OPERATIONS.register("skill", skill)
for operation in SKILL_OPERATIONS:
    OPERATIONS.register(operation, skill)
OPERATIONS.register("image", image)
OPERATIONS.register("document", document)
# Legacy direct callers may omit the descriptor handler; aliases preserve that
# pre-process API while process jobs use the registry-owned handler above.
OPERATIONS.register("image_generate", image)
OPERATIONS.register("image_edit", image)
OPERATIONS.register("document_read", document)
OPERATIONS.register("convert_mp4", convert)
OPERATIONS.register("social_read", _social)


def execute(job):
    _validate_protocol(job)
    result = _validate_result(OPERATIONS.execute(job))
    _validate_usage_journal(job)
    return result


if __name__ == "__main__":
    try:
        job = json.loads(sys.stdin.buffer.read(1024 * 1024))
        with contextlib.redirect_stdout(io.StringIO()):
            result = execute(job)
        print(json.dumps({"protocol": PROTOCOL, "protocolVersion": PROTOCOL_VERSION, "result": result}, ensure_ascii=False))
    except Exception as error:
        code = str(error) if isinstance(error, JobError) else getattr(error, "code", type(error).__name__)
        print(json.dumps({"protocol": PROTOCOL, "protocolVersion": PROTOCOL_VERSION, "error": {"code": code}}))
        sys.exit(1)
